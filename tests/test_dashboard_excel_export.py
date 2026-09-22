"""Кнопка «⬇ Excel с цветами»: те же данные, что в CSV, но с подсветкой. Без настоящей базы."""

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

import dashboard_db

GOOD, WARN, BAD = "00" + dashboard_db._EXCEL_GOOD_FILL, "00" + dashboard_db._EXCEL_WARN_FILL, "00" + dashboard_db._EXCEL_BAD_FILL
NO_FILL = "00000000"


def sheet_from(rows):
    wb = load_workbook(io.BytesIO(dashboard_db._current_to_excel_bytes(pd.DataFrame(rows))))
    return wb.active


def headers(ws):
    return [c.value for c in ws[1]]


def fill(ws, header_label, row=2):
    col = headers(ws).index(header_label) + 1
    return ws.cell(row=row, column=col).fill.fgColor.rgb


def test_headers_match_the_same_russian_labels_as_the_on_screen_table():
    ws = sheet_from([{"our_asin": "B0X", "comp_asin": "B0Y", "our_bsr_delta_24h": -1}])
    assert headers(ws) == ["Наш ASIN", "ASIN конкурента", "Δ BSR наш"]


@pytest.mark.parametrize("delta, expected", [(-5, "good"), (5, "bad"), (0, None), (None, None)])
def test_bsr_delta_is_colored_green_when_it_improved_and_red_when_it_worsened(delta, expected):
    ws = sheet_from([{"our_bsr_delta_24h": delta}])
    got = fill(ws, "Δ BSR наш")
    expected_rgb = {"good": GOOD, "bad": BAD, None: NO_FILL}[expected]
    assert got == expected_rgb


@pytest.mark.parametrize("pct, expected", [(15.0, GOOD), (-15.0, BAD), (0, NO_FILL)])
def test_price_edge_favors_us_when_the_competitor_is_pricier(pct, expected):
    ws = sheet_from([{"price_diff_pct": pct}])
    assert fill(ws, "Разница цен, %") == expected


@pytest.mark.parametrize("stock, expected", [("In Stock", GOOD), ("Low Stock (<5)", WARN), ("Out of Stock", BAD), ("", NO_FILL)])
def test_stock_status_is_colored_by_meaning(stock, expected):
    ws = sheet_from([{"comp_stock": stock}])
    assert fill(ws, "Наличие") == expected


def test_a_decimal_from_postgres_numeric_columns_is_colored_like_a_plain_number():
    from decimal import Decimal
    ws = sheet_from([{"our_bsr_delta_24h": Decimal("-2.50")}])
    assert fill(ws, "Δ BSR наш") == GOOD


def test_a_timezone_aware_updated_at_from_postgres_does_not_crash_excel():
    """updated_at — TIMESTAMPTZ; Excel не понимает часовой пояс в дате и падает, если его не снять."""
    ts = pd.Timestamp("2026-09-22 10:00:00", tz="UTC")
    ws = sheet_from([{"our_asin": "B0X", "comp_asin": "B0Y", "updated_at": ts}])
    assert "Обновлено" in headers(ws)


def test_an_empty_table_does_not_crash():
    columns = ["our_asin", "comp_asin", "our_bsr_delta_24h", "comp_bsr_delta_24h", "price_diff_pct", "comp_stock"]
    ws = sheet_from(pd.DataFrame(columns=columns).to_dict("records"))
    assert headers(ws)
